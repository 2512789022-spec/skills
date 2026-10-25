#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""提取 docx 的版式指纹：页面、样式、标题/正文/题注参数、表格边框、图片尺寸、公式与页码。

用法:
    python inspect_reference.py <参考稿.docx> [--max-paras N]
"""
import argparse
import re
import sys

import docx
from docx.oxml.ns import qn


def cm(v):
    return None if v is None else round(v.cm, 2)


def pt(v):
    return None if v is None else round(v.pt, 1)


def para_fp(p):
    pf = p.paragraph_format
    pPr = p._p.pPr
    lvl = pPr.find(qn('w:outlineLvl')) if pPr is not None else None
    ind = pPr.find(qn('w:ind')) if pPr is not None else None
    sizes = [r.font.size.pt for r in p.runs if r.font.size]
    fonts = set()
    for r in p.runs:
        if r._element.rPr is None:
            continue
        rf = r._element.rPr.find(qn('w:rFonts'))
        if rf is not None:
            fonts.add((rf.get(qn('w:ascii')), rf.get(qn('w:eastAsia'))))
    return {
        'align': str(pf.alignment).split(' ')[0] if pf.alignment is not None else 'inherit',
        'line': pf.line_spacing,
        'before': pt(pf.space_before),
        'after': pt(pf.space_after),
        'indent_cm': cm(pf.first_line_indent),
        'outline': lvl.get(qn('w:val')) if lvl is not None else None,
        'bold': any(r.font.bold for r in p.runs),
        'size': sizes[0] if sizes else None,
        'fonts': sorted(f for f in fonts if any(f)),
        'ind_attrs': dict(ind.attrib) if ind is not None else None,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('path')
    ap.add_argument('--max-paras', type=int, default=400)
    a = ap.parse_args()

    d = docx.Document(a.path)
    s = d.sections[0]
    print('== 页面 ==')
    print(f'  纸张 {cm(s.page_width)} x {cm(s.page_height)} cm；'
          f'边距 上{cm(s.top_margin)} 下{cm(s.bottom_margin)} 左{cm(s.left_margin)} 右{cm(s.right_margin)}')
    text_w = int(s.page_width - s.left_margin - s.right_margin) / 360000
    print(f'  版心宽 {round(text_w, 2)} cm')

    st = d.styles['Normal']
    print('== Normal 样式 ==')
    print(f'  size={pt(st.font.size)} alignment={st.paragraph_format.alignment}')

    body = d.element.body
    print('== 规模 ==')
    print(f'  段落 {len(d.paragraphs)}；表格 {len(d.tables)}；'
          f'OMML 公式 {len(body.findall(".//" + qn("m:oMath")))}；'
          f'图片 inline {len(body.findall(".//" + qn("wp:inline")))} '
          f'anchor {len(body.findall(".//" + qn("wp:anchor")))}')

    print('== 段落清单（非空） ==')
    for i, p in enumerate(d.paragraphs[:a.max_paras]):
        t = p.text.strip()
        imgs = len(p._p.findall('.//' + qn('a:blip')))
        if not t and not imgs:
            continue
        tag = '[IMG]' if imgs else ''
        print(f'  P{i:<3} {para_fp(p)} {tag} {t[:50]!r}')

    print('== 图片尺寸 ==')
    for i, p in enumerate(d.paragraphs):
        for e in p._p.findall('.//' + qn('wp:extent')):
            cx, cy = int(e.get('cx')) / 360000, int(e.get('cy')) / 360000
            print(f'  P{i:<3} {cx:.2f} x {cy:.2f} cm')

    print('== 表格 ==')
    for ti, t in enumerate(d.tables):
        tb = t._tbl.tblPr.find(qn('w:tblBorders'))
        desc = ''
        if tb is not None:
            desc = ' '.join(f'{c.tag.split("}")[1]}:{c.get(qn("w:val"))}/{c.get(qn("w:sz"))}' for c in tb)
        print(f'  T{ti} {len(t.rows)}x{len(t.columns)} imgs={len(t._tbl.findall(".//" + qn("a:blip")))} borders[{desc}]')
        print('      ', [c.text.strip()[:12] for c in t.rows[0].cells])

    print('== 页脚 ==')
    for i, sec in enumerate(d.sections):
        ft = sec.footer._element
        txt = ''.join(x.text or '' for x in ft.iter(qn('w:t')))
        instr = [x.text for x in ft.iter(qn('w:instrText'))]
        print(f'  节{i} 文本{txt!r} 域{instr}')

    print('== 题注编号 ==')
    caps = [p.text.strip() for p in d.paragraphs if re.match(r'^[图表]\s?\d+', p.text.strip())]
    for c in caps:
        print('  ', c[:60])
    print(f'  共 {len(caps)} 条')


if __name__ == '__main__':
    sys.exit(main())
